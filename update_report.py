import docx
import sys
import datetime

def append_update(update_text):
    try:
        doc = docx.Document('Aura_Project_Progress.docx')
    except Exception as e:
        print("Could not open document:", e)
        return

    # Check if 'Changelog' heading exists, if not, add it.
    changelog_exists = any("Changelog & Recent Updates" in p.text for p in doc.paragraphs)
    
    if not changelog_exists:
        doc.add_heading('6. Changelog & Recent Updates', level=1)
    
    # Add the update with the current date/time
    date_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f"[{date_str}] ").bold = True
    p.add_run(update_text)
    
    doc.save('Aura_Project_Progress.docx')
    print("Document successfully updated!")

if __name__ == '__main__':
    if len(sys.argv) > 1:
        update_text = " ".join(sys.argv[1:])
        append_update(update_text)
    else:
        print("Please provide the update text.")
